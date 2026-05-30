#!/usr/bin/env python3
"""
design_code.py — Tokens de design centralizados do livro (Design Code v1.0)

Fonte única de verdade para a identidade visual da obra:
  - Cores (paleta editorial)
  - Fontes (serifada de corpo / sem-serifa de etiquetas)
  - Escala tipográfica do corpo e dos títulos estruturais
  - Geometria da página A5 com margens espelhadas

Também concentra dois helpers de baixo nível reutilizados por
md_to_docx.py e gerar_livro_pdf.py (margens A5 e campo PAGE), eliminando
a duplicação que antes existia entre os dois scripts.

Este módulo depende APENAS de python-docx — nunca importe daqui os
módulos de geração, para não criar dependência circular.
"""

from docx.shared import Pt, Cm, RGBColor
from docx.enum.section import WD_ORIENT
from docx.oxml.ns import qn, nsdecls
from docx.oxml import parse_xml


# ═══════════════════════════════════════════════════════
# NAMESPACES (Word 2010+)
# ═══════════════════════════════════════════════════════

# Ligaduras OpenType e efeitos de texto avançados
W14_NS = 'http://schemas.microsoft.com/office/word/2010/wordml'
MC_NS = 'http://schemas.openxmlformats.org/markup-compatibility/2006'


# ═══════════════════════════════════════════════════════
# CORES DO DESIGN CODE
# ═══════════════════════════════════════════════════════

COLORS = {
    "body_text":     "201B16",
    "title_text":    "2F2923",
    "heading_text":  "3A3128",
    "muted_text":    "6B6256",
    "gold_accent":   "B78B37",
    "gold_dark":     "8A6A2A",
    "rule_light":    "D7CDBB",
    "box_border":    "D2C9B8",
    "box_atencao_bg":    "F7F1E6",
    "box_juris_bg":      "EEF3F7",
    "box_pratica_bg":    "F0F4ED",
    "box_atencao_accent":  "B78B37",
    "box_juris_accent":    "4D6F8A",
    "box_pratica_accent":  "6E8563",
    "box_quadro_bg":       "F5F3EF",
    "box_quadro_accent":   "6B6256",
    "box_quadro_border":   "A09882",
}


def hex_to_rgb(hex_str):
    """Converte string hexadecimal (RRGGBB) em RGBColor do python-docx."""
    return RGBColor(int(hex_str[:2], 16), int(hex_str[2:4], 16), int(hex_str[4:6], 16))


# ═══════════════════════════════════════════════════════
# FONTES
# ═══════════════════════════════════════════════════════

FONT_SERIF = "EB Garamond"   # corpo, títulos e a maior parte do texto
FONT_SANS = "Noto Sans"      # etiquetas (CAPÍTULO N, PARTE, rótulos de box)


# ═══════════════════════════════════════════════════════
# ESCALA TIPOGRÁFICA (corpo + títulos estruturais)
# ═══════════════════════════════════════════════════════
# Tamanhos compartilhados pelos dois geradores. Mantê-los aqui evita
# divergência entre md_to_docx.py (capítulo isolado) e gerar_livro_pdf.py
# (livro unificado), que antes repetiam os mesmos valores.

SIZE_BODY = Pt(10.7)         # texto corrido
SIZE_SECTION = Pt(15.5)      # título de seção  (### X.Y)
SIZE_SUBSECTION = Pt(12.5)   # título de subseção (#### X.Y.Z)


# ═══════════════════════════════════════════════════════
# GEOMETRIA DA PÁGINA A5 (margens espelhadas)
# ═══════════════════════════════════════════════════════

MARGINS_A5 = {
    "page_width_cm": 14.8,
    "page_height_cm": 21.0,
    "top_cm": 1.65,
    "bottom_cm": 1.70,
    "left_cm": 2.05,    # interna (lombada)
    "right_cm": 1.55,   # externa
    "header_cm": 0.80,
    "footer_cm": 0.75,
}


def apply_a5_margins(section):
    """Aplica dimensões e margens A5 espelhadas a uma seção.

    Define largura/altura, as quatro margens, as distâncias de
    cabeçalho/rodapé e ativa mirrorMargins no pgMar da seção. Os ajustes
    de nível de documento (mirror em settings, hifenização, w14) ficam a
    cargo de quem chama, pois variam entre os dois geradores.
    """
    section.page_width = Cm(MARGINS_A5["page_width_cm"])
    section.page_height = Cm(MARGINS_A5["page_height_cm"])
    section.orientation = WD_ORIENT.PORTRAIT
    section.top_margin = Cm(MARGINS_A5["top_cm"])
    section.bottom_margin = Cm(MARGINS_A5["bottom_cm"])
    section.left_margin = Cm(MARGINS_A5["left_cm"])
    section.right_margin = Cm(MARGINS_A5["right_cm"])
    section.header_distance = Cm(MARGINS_A5["header_cm"])
    section.footer_distance = Cm(MARGINS_A5["footer_cm"])

    sectPr = section._sectPr
    pgMar = sectPr.find(qn('w:pgMar'))
    if pgMar is not None:
        pgMar.set(qn('w:mirrorMargins'), '1')


# ═══════════════════════════════════════════════════════
# CAMPO PAGE (rodapés)
# ═══════════════════════════════════════════════════════

def add_page_field(paragraph, font_name=FONT_SERIF, size=Pt(9), color="6B6256"):
    """Insere um campo PAGE num parágrafo (helper para rodapés)."""
    run1 = paragraph.add_run()
    run1.font.name = font_name
    run1.font.size = size
    run1.font.color.rgb = hex_to_rgb(color)
    fld_begin = parse_xml(f'<w:fldChar {nsdecls("w")} w:fldCharType="begin"/>')
    run1._element.append(fld_begin)

    run2 = paragraph.add_run()
    run2.font.name = font_name
    run2.font.size = size
    run2.font.color.rgb = hex_to_rgb(color)
    instr = parse_xml(f'<w:instrText {nsdecls("w")} xml:space="preserve"> PAGE </w:instrText>')
    run2._element.append(instr)

    run3 = paragraph.add_run()
    run3.font.name = font_name
    run3.font.size = size
    fld_end = parse_xml(f'<w:fldChar {nsdecls("w")} w:fldCharType="end"/>')
    run3._element.append(fld_end)
