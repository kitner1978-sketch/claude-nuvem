#!/usr/bin/env python3
"""
verificar_jurisprudencia.py — Cruza as citações de Temas vinculantes (STJ, STF
e TNU) feitas nos 23 capítulos do livro contra os dados oficiais baixados:

  - STJ: jurisprudencia/stj/precedentes/  (precedentes qualificados — CKAN)
  - STF: jurisprudencia/stf/temas_repercussao_geral_stf.xlsx  (Corte Aberta)
  - TNU: jurisprudencia/tnu/temas_representativos_tnu.csv  (scraping CJF)

Detecta, por tribunal:
  - ⛔ Tema citado que NÃO existe (fabricação / número errado / tribunal trocado)
  - ⚠️ Tema cujo número existe mas a TESE é de assunto alheio ao previdenciário
        (pega citações com o número certo de um Tema errado)
  - ✅ Tema confirmado, com tese/relator/data oficiais para conferência

Gera: output/verificacao_jurisprudencia.md
"""
import csv, os, re
from collections import defaultdict

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
STJ = os.path.join(ROOT, "jurisprudencia", "stj", "precedentes")
STF_XLSX = os.path.join(ROOT, "jurisprudencia", "stf", "temas_repercussao_geral_stf.xlsx")
TNU_CSV = os.path.join(ROOT, "jurisprudencia", "tnu", "temas_representativos_tnu.csv")
CAPS = os.path.join(ROOT, "output", "rascunhos")
OUT = os.path.join(ROOT, "output", "verificacao_jurisprudencia.md")

n = lambda s: (s or "").strip()
dig = lambda s: re.sub(r"\D", "", s or "")

RX = {
    "STJ": re.compile(r"Tema\s+(\d[\d.]*)\s*(?:/|\bdo\b)\s*STJ", re.I),
    "STF": re.compile(r"Tema\s+(\d[\d.]*)\s*(?:/|\bdo\b)\s*STF", re.I),
    "TNU": re.compile(r"Tema\s+(\d[\d.]*)\s*(?:/|\b(?:da|do)\b)\s*TNU", re.I),
}

PREV_KW = ["previd", "segurad", "aposentad", "auxilio", "auxílio", "pensao", "pensão",
           "beneficio", "benefício", "inss", "rgps", "loas", "bpc", "carenc", "carênc",
           "incapac", "deficien", "rural", "reclus", "maternidade", "acidente",
           "seguridade", "assistência social", "assistencia social", "contribui",
           "8.213", "8213", "salário-de-contribuição", "salario-de-contribuicao"]
ALIEN_KW = ["comissão de permanência", "comissao de permanencia", "imposto sobre serviç",
            "tributári", "tributari", " iss ", "icms", "consumidor",
            "alienação fiduciária", "arrendamento mercantil", "locaç", "penal",
            "criminal", "bancári", "bancari", "cartão de crédito", "execução fiscal",
            "ipva", "iptu", "condomínio", "marca", "patente", "ambiental", "improbidade"]


def assunto_suspeito(texto):
    t = (texto or "").lower()
    return (not any(k in t for k in PREV_KW)) and any(k in t for k in ALIEN_KW)


# ── Carregamento das três bases ──────────────────────────────────────────────

def load_stj():
    rd = lambda fn: list(csv.DictReader(open(os.path.join(STJ, fn), encoding="utf-8", errors="replace")))
    temas, procs = rd("temas.csv"), rd("processos.csv")
    pidx = defaultdict(list)
    for p in procs:
        if "Tema" in n(p.get("tipoPrecedente")):
            pidx[dig(p.get("numeroPrecedente"))].append(p)
    idx = {}
    for t in temas:
        if "Tema" in n(t.get("tipoPrecedente")):
            num = dig(t.get("numeroPrecedente"))
            ps = pidx.get(num, [])
            lead = next((p for p in ps if n(p.get("leadingCase")) in ("S", "Sim")), ps[0] if ps else {})
            idx[num] = {"tese": n(t.get("teseFirmada")), "data": n(t.get("dataJulgamento")),
                        "rel": n(lead.get("ministroRelator")), "proc": n(lead.get("Processo")),
                        "sit": n(t.get("situacao")), "assunto": n(t.get("Assuntos"))}
    return idx


def load_stf():
    from openpyxl import load_workbook
    ws = load_workbook(STF_XLSX, read_only=True, data_only=True)["Sheet1"]
    it = ws.iter_rows(values_only=True)
    hdr = [n(str(c)) for c in next(it)]
    h = {name: i for i, name in enumerate(hdr)}
    g = lambda r, k: n(str(r[h[k]])) if r[h[k]] is not None else ""
    idx = {}
    for r in it:
        if not r or r[h["Número tema"]] in (None, ""):
            continue
        num = dig(g(r, "Número tema"))
        tese = g(r, "Tese").replace("_x000D_", " ")
        idx[num] = {"tese": tese, "data": g(r, "Data julgamento tema")[:10],
                    "rel": g(r, "Relator atual"), "proc": g(r, "Processo paradigma"),
                    "sit": g(r, "Situação julgamento tema"), "assunto": g(r, "Ramos do Direito")}
    return idx


def load_tnu():
    idx = {}
    for r in csv.DictReader(open(TNU_CSV, encoding="utf-8", errors="replace")):
        num = dig(r.get("numero"))
        if not num:
            continue
        idx[num] = {"tese": n(r.get("tese")), "data": n(r.get("Julgado em")),
                    "rel": n(r.get("Relator (a)")), "proc": n(r.get("processo")),
                    "sit": n(r.get("situacao")), "assunto": n(r.get("ramo"))}
    return idx


def scan():
    cites = {trib: defaultdict(set) for trib in RX}
    files = sorted(f for f in os.listdir(CAPS) if re.match(r"cap_\d+_rascunho\.md$", f))
    for fn in files:
        cap = re.search(r"cap_(\d+)_", fn).group(1)
        text = open(os.path.join(CAPS, fn), encoding="utf-8", errors="replace").read()
        for trib, rx in RX.items():
            for m in rx.finditer(text):
                cites[trib][dig(m.group(1))].add(cap)
    return cites, files


# ── Relatório ────────────────────────────────────────────────────────────────

def main():
    bases = {"STJ": load_stj(), "STF": load_stf(), "TNU": load_tnu()}
    cites, files = scan()

    achados = {trib: {"ok": [], "nf": [], "susp": []} for trib in RX}
    for trib in RX:
        idx = bases[trib]
        for num in sorted(cites[trib], key=lambda x: int(x) if x.isdigit() else 0):
            caps = ", ".join(sorted(cites[trib][num], key=int))
            rec = idx.get(num)
            if not rec:
                achados[trib]["nf"].append((num, caps))
                continue
            rec = {**rec, "num": num, "caps": caps}
            achados[trib]["ok"].append(rec)
            if assunto_suspeito(rec["tese"] + " " + rec.get("assunto", "")):
                achados[trib]["susp"].append(rec)

    L = ["# Verificação de Jurisprudência — Temas vinculantes citados no livro\n"]
    L.append("> Fontes oficiais: **STJ** (precedentes qualificados), **STF** (Temas de "
             "Repercussão Geral), **TNU** (Temas Representativos).\n")
    L.append(f"- Capítulos varridos: {len(files)}")
    for trib in RX:
        a = achados[trib]
        L.append(f"- **{trib}**: {len(cites[trib])} citados | ✅ {len(a['ok'])} | "
                 f"⛔ {len(a['nf'])} não encontrados | ⚠️ {len(a['susp'])} assunto suspeito")
    L.append("")

    # Não encontrados
    nf_any = any(achados[t]["nf"] for t in RX)
    if nf_any:
        L.append("## ⛔ Citados mas NÃO encontrados na base oficial\n")
        for trib in RX:
            for num, caps in achados[trib]["nf"]:
                L.append(f"- **Tema {num}/{trib}** — caps. {caps}")
        L.append("")

    # Assunto suspeito
    susp_any = any(achados[t]["susp"] for t in RX)
    if susp_any:
        L.append("## ⚠️ Número existe, mas a TESE é de assunto alheio ao previdenciário\n")
        L.append("(provável citação do número errado — conferir)\n")
        for trib in RX:
            for r in achados[trib]["susp"]:
                L.append(f"- **Tema {r['num']}/{trib}** (caps. {r['caps']}) — "
                         f"assunto: {r.get('assunto','')[:50]} — tese: {r['tese'][:90]}…")
        L.append("")

    # Confirmados por tribunal
    for trib in RX:
        ok = achados[trib]["ok"]
        if not ok:
            continue
        L.append(f"## ✅ {trib} — Temas confirmados ({len(ok)})\n")
        L.append("| Tema | Caps | Julgamento | Relator | Processo | Situação |")
        L.append("|------|------|-----------|---------|----------|----------|")
        for r in ok:
            flag = " ⚠️" if r in achados[trib]["susp"] else ""
            L.append(f"| {r['num']}{flag} | {r['caps']} | {r['data']} | {r['rel']} | "
                     f"{r['proc']} | {r['sit']} |")
        L.append("")

    open(OUT, "w", encoding="utf-8").write("\n".join(L))
    tot_c = sum(len(cites[t]) for t in RX)
    tot_nf = sum(len(achados[t]["nf"]) for t in RX)
    tot_su = sum(len(achados[t]["susp"]) for t in RX)
    print(f"Relatório: {OUT}")
    for trib in RX:
        a = achados[trib]
        print(f"  {trib}: {len(cites[trib])} citados | ok {len(a['ok'])} | "
              f"nao-encontrados {len(a['nf'])} | suspeitos {len(a['susp'])}")
    print(f"TOTAL: {tot_c} Temas citados | {tot_nf} nao encontrados | {tot_su} assunto suspeito")
    for trib in RX:
        if achados[trib]["nf"]:
            print(f"  [{trib}] NAO ENCONTRADOS:", ", ".join(x[0] for x in achados[trib]["nf"]))
        if achados[trib]["susp"]:
            print(f"  [{trib}] SUSPEITOS:", ", ".join(r["num"] for r in achados[trib]["susp"]))


if __name__ == "__main__":
    main()
